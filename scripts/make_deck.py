"""Deck del pitch (issue #16): docs/pitch/deck.pptx + docs/pitch/deck.md desde una sola fuente.

    PYTHONPATH=src python scripts/make_charts.py      # primero (o --from-cache sin datos)
    PYTHONPATH=src python scripts/make_deck.py        # 6 slides con notas del orador

- Los números del problema salen de docs/pitch/chart_data.json (agregados de make_charts.py).
- Las métricas de evaluación/capacidad salen de docs/pitch/metrics.yaml; si faltan, la slide muestra
  [[PENDIENTE: ...]] y el script las lista al final (no se inventa ningún número).
- Escenas de la slide 3: clientes, mensajes y cifras de docs/demo/demo_customers.json (el mismo guion que la demo
  y el video); el conteo de tests sale de `pytest --collect-only -q`.
- Capturas de la demo: docs/pitch/screens/escena{1,2,3}.png; si no existen, se dibuja un marcador.
"""
from __future__ import annotations

import glob
import json
import os
import re

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PITCH = os.path.join(ROOT, "docs", "pitch")
CHARTS = os.path.join(PITCH, "charts")
REPO = "github.com/Mr-Percival/factored-hackathon-2026-disputas"

INK, INK_2, INK_3 = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x8A, 0x89, 0x84)
BLUE, ORANGE = RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34)
PENDING = RGBColor(0xB5, 0x4A, 0x1A)  # naranja oscuro legible sobre blanco para [[PENDIENTE]]
WHITE, CARD, TINT = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xF4, 0xF4, 0xF2), RGBColor(0xEA, 0xF2, 0xFC)
FONT = "Calibri"

pending: list[str] = []


# ------------------------------------------------------------------ insumos
def load():
    data = json.load(open(os.path.join(PITCH, "chart_data.json"), encoding="utf-8"))
    metrics = yaml.safe_load(open(os.path.join(PITCH, "metrics.yaml"), encoding="utf-8"))
    try:  # conteo real (incluye parametrizados); si pytest no está, se cuentan las funciones test_
        import subprocess
        import sys
        r = subprocess.run([sys.executable, "-m", "pytest", "--collect-only", "-q"], cwd=ROOT,
                           capture_output=True, text=True, timeout=120)
        hit = re.search(r"(\d+) tests? collected", r.stdout)
        if hit:
            return data, metrics, int(hit.group(1))
    except Exception:  # noqa: BLE001
        pass
    n_tests = sum(len(re.findall(r"^\s*def test_", open(f, encoding="utf-8").read(), re.M))
                  for f in glob.glob(os.path.join(ROOT, "tests", "test_*.py")))
    return data, metrics, n_tests


def m(metrics: dict, key: str) -> str:
    """Valor de metrics.yaml o marcador [[PENDIENTE]] (y lo registra)."""
    item = metrics[key]
    if item.get("value") not in (None, ""):
        return str(item["value"])
    tag = f"[[PENDIENTE: {item['desc']} — {item['source']}]]"
    if tag not in pending:
        pending.append(tag)
    return f"[[PENDIENTE: {item['source'].split(' ')[0]}]]"


def demo_scenes() -> dict:
    """Escenarios del guion real (docs/demo/demo_customers.json) por id."""
    d = json.load(open(os.path.join(ROOT, "docs", "demo", "demo_customers.json"), encoding="utf-8"))
    return {sc["id"]: sc for sc in d["scenarios"]}


def money(x: float, cur: str) -> str:
    return f"{x:,.2f} {cur}"


def facts(data: dict) -> dict:
    total = sum(r["n"] for r in data["demand"])
    q = next(r for r in data["demand"] if r["category"] == "Queja")
    tr = next(r for r in data["demand"] if r["category"] == "Transaccional")
    pqr_total = sum(r["n"] for r in data["pqr"])
    disp = [r for r in data["pqr"] if r["subcategory"] in ("Cargo no reconocido", "Cobro indebido")]
    dn = sum(r["n"] for r in disp)
    qd = data["quality"]
    return {
        "queja_share": f"{100 * q['n'] / total:.0f}%", "queja_fcr": f"{100 * q['fcr']:.1f}%",
        "trx_fcr": f"{100 * tr['fcr']:.1f}%", "queja_followup": f"{100 * q['followup']:.0f}%",
        "dur_x": f"{q['median_duration_s'] / tr['median_duration_s']:.1f} veces más",
        "disp_share": f"{100 * dn / pqr_total:.1f}%", "disp_n": f"{dn:,}".replace(",", "."),
        "sla": f"{100 * sum(r['sla_breach'] * r['n'] for r in disp) / dn:.0f}%",
        "med_days": f"{max(r['median_days'] for r in disp):.0f}",
        "texts": str(qd["distinct_customer_texts"]), "fraud_null": f"{100 * qd['nulls']['fraud_score']:.0f}%",
        "ids_changed": f"{100 * qd['snapshot']['customers_ids_only_in_one'] / qd['snapshot']['customers_total']:.0f}%",
    }


# ------------------------------------------------------------------ primitivas de dibujo
def text(slide, x, y, w, h, content, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False):
    """content: str (líneas con \\n) o lista de párrafos [(texto, {opciones})]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [(line, {}) for line in content.split("\n")]
    for i, (t, o) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.space_after = Pt(o.get("space", 4))
        if o.get("mark"):  # marcador (✓ / ○) en una fuente de símbolos, en azul
            mk = p.add_run()
            mk.text = o["mark"] + "  "
            mk.font.name, mk.font.size, mk.font.bold = "Segoe UI Symbol", Pt(o.get("size", size)), True
            mk.font.color.rgb = BLUE
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = FONT
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.italic = o.get("italic", italic)
        c = o.get("color", color)
        f.color.rgb = PENDING if "[[PENDIENTE" in t else c
    return tb


def box(slide, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(2)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def label_in(shape, t, size=15, color=WHITE, bold=True):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size if i == 0 else size - 3), bold if i == 0 else False
        r.font.color.rgb = color


def header(slide, n, kicker, title):
    circ = box(slide, 0.6, 0.45, 0.42, 0.42, fill=BLUE, shape=MSO_SHAPE.OVAL)
    label_in(circ, str(n), size=14)
    text(slide, 1.15, 0.47, 10, 0.4, kicker, size=13, color=INK_3, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 0.6, 1.0, 12.1, 0.9, title, size=32, bold=True)


def footer(slide, note):
    text(slide, 0.6, 7.0, 12.1, 0.3, note, size=10, color=INK_3)


def image(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    return s


# ------------------------------------------------------------------ contenido (fuente única pptx + md)
def build(data, metrics, n_tests):
    f = facts(data)
    titles = json.load(open(os.path.join(CHARTS, "titles.json"), encoding="utf-8"))
    demo_url = m(metrics, "demo_url")
    sc = demo_scenes()
    n, a_, fr, inj = sc["normal_mx"], sc["ambiguous_co"], sc["fraud_pt"], sc["injection"]
    nt, ft = n["transaction"], fr["transaction"]
    return [
        {  # 1 ----------------------------------------------------------------
            "kicker": "El problema, con datos",
            "title": "Disputar un cargo es el peor momento del servicio",
            "chart": "01_demanda_quejas.png",
            "tiles": [(f["queja_fcr"], f"de las quejas se resuelve al primer contacto (vs {f['trx_fcr']} transaccional)"),
                      (f["disp_share"], f"de los PQR son «cargo no reconocido» o «cobro indebido» ({f['disp_n']} casos)"),
                      (f"{f['med_days']} días", f"mediana de resolución; {f['sla']} incumple el SLA")],
            "foot": "Datos sintéticos del organizador (LATAM Bank) · 686.296 interacciones y 67.095 PQR · reports/EDA_findings.md",
            "md_extra": [f"Gráfico: _{titles['01_demanda_quejas']}_", f"Gráfico de respaldo: _{titles['02_pqr_subcategorias']}_ "
                         "(`charts/02_pqr_subcategorias.png`)"],
            "notes": (
                f"Empezamos por los datos, no por la tecnología. En el call center, las quejas son el {f['queja_share']} de los "
                f"contactos pero solo el {f['queja_fcr']} se resuelve al primer contacto, contra {f['trx_fcr']} en consultas "
                f"transaccionales; el {f['queja_followup']} necesita seguimiento y duran {f['dur_x']}. En PQR, cargo no "
                f"reconocido y cobro indebido son el {f['disp_share']}, con un {f['sla']} fuera de SLA y {f['med_days']} días de "
                "mediana. Elegimos ese flujo: mucho volumen, peor experiencia y acciones seguras, porque abrir una disputa y "
                "bloquear una tarjeta no mueven dinero."),
        },
        {  # 2 ----------------------------------------------------------------
            "kicker": "Qué construimos",
            "title": "El modelo interpreta; el código decide y actúa",
            "diagram": True,
            "layers": [("Datos", "Contrato por tabla, cuarentena, upsert incremental y linaje por corrida"),
                       ("Herramientas", "Única puerta a los datos: el cliente sale del token de sesión, nunca del texto"),
                       ("Política", "policy.yaml versionado: reglas con ID (D01…H04) que explican cada decisión"),
                       ("Agente", "LangGraph con transiciones declaradas; 19 caminos verificados, lo no declarado falla cerrado")],
            "foot": "Detalle: docs/architecture.md · ADR 0002 (grafo determinista) · ADR 0009 (router aprendido)",
            "md_extra": ["Diagrama: guard → entender (router aprendido local: intención; reglas: datos y consentimiento) → "
                         "buscar → decidir (policy.yaml) → actuar con confirmación y verificación / handoff → responder "
                         "(plantillas es/pt; writer LLM opcional, apagado). Render de respaldo: `charts/arq_4_capas.png`, "
                         "`charts/arq_grafo.png`."],
            "notes": (
                "Un solo agente en LangGraph con pasos explícitos. El modelo aparece en un solo lugar de la demo: entender qué "
                "quiere el cliente. Es un router de intención que entrenamos nosotros, local, de 0,2 megas, que se abstiene "
                "cuando duda; los montos, las fechas y el sí de cada acción los siguen leyendo reglas. Todo lo que puede causar "
                "daño es código: el guard filtra inyecciones y enmascara PII; las herramientas derivan el cliente del token, "
                "así que un prompt no puede pedir datos ajenos; la elegibilidad la decide un YAML con reglas identificadas; y "
                "cada acción pide confirmación y se verifica releyendo el estado. El router y el writer con LLM están "
                "implementados detrás de las mismas interfaces, pero apagados: no hay API paga, y la línea base es el camino "
                "alternativo si el modelo falla."),
        },
        {  # 3 ----------------------------------------------------------------
            "kicker": "Demo en 3 escenas · clientes reales del dataset",
            "title": "Un solo flujo: resuelve, pregunta o deriva — nunca adivina",
            "scenes": [
                (f"Normal ({n['locale']})", f"«{n['turns'][0]['say']}»",
                 f"1 cargo ({money(nt['amount'], nt['currency'])}, {nt['merchant']}). Ofrece disputa y bloqueo por "
                 f"separado; con «{n['turns'][1]['say']}» abre la disputa y bloquea la tarjeta, ambas verificadas."),
                (f"Ambiguo ({a_['locale']})", f"«{a_['turns'][0]['say']}»",
                 f"{len(a_['candidate_ids'])} candidatos: pregunta cuál y solo avanza cuando el cliente elige "
                 f"(«{a_['turns'][1]['say']}»)."),
                (f"Handoff en pt + ataque", f"«{fr['turns'][0]['say']}»",
                 f"fraud_score {ft['fraud_score']:g} ≥ 30 → regla H03 → analista con el paquete completo. "
                 f"«{' '.join(inj['turns'][0]['say'].split()[:4])}…» → bloqueado por el guard."),
            ],
            "foot": f"Demo en vivo: {demo_url} · clientes de docs/demo/demo_customers.json (verificados 14/14) · guion: docs/demo-script.md",
            "notes": (
                f"Tres conversaciones del mismo flujo, con clientes reales del dataset del organizador. Primera, en México: "
                f"el cliente no reconoce un cargo de {nt['amount']:.0f} dólares en {nt['merchant']}; hay un único candidato, "
                f"la política lo permite y el agente ofrece la disputa y el bloqueo por separado. Con «{n['turns'][1]['say']}» "
                "abre la disputa y bloquea la tarjeta; solo dice 'listo' si la verificación lo confirma. Segunda, en Colombia: "
                f"{a_['turns'][0]['say'].split('de ')[1].split(' que')[0]} coincide con dos cargos, así que pregunta en vez de "
                "adivinar. Tercera, en portugués: un cargo con score de fraude "
                f"{ft['fraud_score']:g} activa la regla H03 y pasa a un analista con la traza, los hechos verificados y la regla. "
                "Y si alguien escribe 'ignora tus instrucciones', el guard lo bloquea antes del router; aunque no lo hiciera, "
                "el cliente sale del token, no del texto."),
        },
        {  # 4 ----------------------------------------------------------------
            "kicker": "Cómo sabemos que funciona",
            "title": "Evaluación honesta: contra líneas base, en casos reservados",
            "metrics": [
                ("Agente de la demo", m(metrics, "eval_agent"), m(metrics, "eval_agent_detail"),
                 f"{m(metrics, 'eval_rules')} · {m(metrics, 'eval_baseline')}"),
                ("Componente aprendido", m(metrics, "ml_router"), m(metrics, "ml_router_detail"),
                 "Router de intención local vs reglas · corpus propio, split por autor (n=306)"),
                ("Inyección y red team", "0 daños", m(metrics, "inj_harm"),
                 f"Red team adaptativo: {m(metrics, 'redteam')}; el daño real se corrigió en la raíz"),
                ("Equidad por dialecto", m(metrics, "fairness"), "grupos con la misma decisión MX · CO · AR · pt-BR",
                 "No enrutamos por acento: el EDA no muestra efecto (FCR 76.7% vs 76.6%)"),
            ],
            "foot": (f"Test congelado: {m(metrics, 'eval_cases')} (segunda medición, fuga leve declarada) · "
                     f"{n_tests} tests en CI, sin datos ni API keys"),
            "notes": (
                "No entrenamos sobre las transcripciones del dataset: solo hay 42 textos de cliente distintos. Medimos donde "
                "hay etiquetas válidas. Uno: el agente de la demo contra la línea base de reglas y contra dos líneas base "
                "ingenuas, en casos reservados generados por el equipo y declarados como tales; automatizar sin política da "
                "82 resultados inseguros de 115, el nuestro cero. Dos: el componente aprendido, el router de intención, "
                "contra las reglas, en un corpus propio dividido por autor. Tres: seguridad; con datasets públicos, "
                f"{m(metrics, 'inj_detection')}, pero lo que importa es el daño, que medimos aparte: cero, aun con el guard "
                "apagado. Y un red team propio de tres rondas sí encontró daño real, que corregimos en la raíz. Cuatro: "
                "equidad; el mismo caso en dialectos distintos debe dar la misma decisión. Somos explícitos: el test ya se "
                "midió dos veces y declaramos una fuga leve."),
        },
        {  # 5 ----------------------------------------------------------------
            "kicker": "Datos e ingeniería que lo sostienen",
            "title": "Decisiones tomadas por los datos, no a pesar de ellos",
            "chart": "04_umbral_fraude.png",
            "limits": [(f["texts"], "textos de cliente distintos en 171k transcripciones", "no entrenamos intención con ellos"),
                       (f["fraud_null"], "de transacciones sin fraud_score", "riesgo residual documentado; H02/H04 y el analista"),
                       (f["ids_changed"], "de IDs de cliente cambian entre las 2 entregas", "huella por entrega + upsert idempotente + linaje")],
            "foot": "reports/fraud_threshold.md · ADR 0006 · charts/03_calidad_datos.png (6 límites → 6 decisiones)",
            "md_extra": [f"Gráfico: _{titles['04_umbral_fraude']}_",
                         f"Respaldo: _{titles['03_calidad_datos']}_ (`charts/03_calidad_datos.png`)"],
            "notes": (
                "La política no se escribió a ojo. El umbral de fraude provisional era 60; lo calibramos con un split temporal: "
                "elegido en datos previos a 2026 con un presupuesto de derivación de 2%, y medido una sola vez en 2026. Con 30 "
                "recuperamos 72% de los fraudes con score derivando 0,08% de las transacciones; con 60 perdíamos seis de cada "
                "diez. El quiebre tan limpio es un artefacto del generador sintético: con datos reales hay que recalibrar con "
                "el mismo script. Y cada límite de los datos se reporta y cambió una decisión."),
        },
        {  # 6 ----------------------------------------------------------------
            "kicker": "Ruta a producción",
            "title": "Qué está listo en el sandbox y qué falta para producción",
            "ready": ["Permisos por sesión: token con expiración, sin IDs en el texto",
                      "Política versionada con reglas identificadas y umbral calibrado",
                      "Consentimiento por acción, acciones verificadas, handoff versionado",
                      "Router aprendido local con fallback visible a reglas",
                      "Red team de 3 rondas e invariantes con Hypothesis, en CI",
                      "Auditoría encadenada · dependencias con hash · CI → GHCR → Render",
                      f"Demo desplegada: {demo_url}"],
            "missing": ["IdP real (OIDC) con MFA y step-up antes de actuar",
                        "Estado compartido (checkpointer, sesiones) para escalar",
                        "Auditoría WORM, claves en KMS, gateway con WAF",
                        "Evaluar el LLM (hoy apagado) y un test v3 escrito por otra persona",
                        "Calificaciones humanas de los handoffs (kit listo)",
                        "Recalibrar el umbral con datos reales",
                        f"Capacidad hoy: {m(metrics, 'capacity')}"],
            "foot": f"{REPO} · docs/threat-model.md · docs/adr/",
            "notes": (
                "Cerramos con lo que falta, sin maquillaje. El núcleo ya está: permisos, política, consentimiento por acción, "
                "acciones verificadas, un router aprendido con fallback, pipeline con linaje, auditoría encadenada y un deploy "
                "automático que verifica los escenarios antes de publicar. Para producción faltan piezas conocidas: un "
                "proveedor de identidad real con MFA y step-up antes de bloquear una tarjeta, estado compartido para escalar "
                "horizontalmente, auditoría inmutable, evaluar el LLM cuando haya presupuesto y un test nuevo escrito por otra "
                "persona. Y con datos reales, recalibrar el umbral. La arquitectura no cambia: cambian los adaptadores."),
        },
    ]


# ------------------------------------------------------------------ render pptx
def render_pptx(slides, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, sd in enumerate(slides, 1):
        s = new_slide(prs)
        header(s, i, sd["kicker"], sd["title"])
        if i == 1:
            image(s, os.path.join(CHARTS, sd["chart"]), 0.5, 2.05, w=8.6)
            for k, (big, small) in enumerate(sd["tiles"]):
                y = 2.1 + k * 1.58
                box(s, 9.4, y, 3.35, 1.42, fill=TINT)
                text(s, 9.6, y + 0.12, 3.0, 0.62, big, size=34, bold=True, color=BLUE)
                text(s, 9.6, y + 0.78, 3.0, 0.6, small, size=12.5, color=INK)
        elif i == 2:
            nodes = [("Guard", "inyección · PII", "code"), ("Entender", "intención: modelo local", "llm"),
                     ("Buscar", "transacción / tarjeta", "code"), ("Decidir", "policy.yaml", "code"),
                     ("Actuar / Handoff", "confirmado + verificado", "code"), ("Responder", "plantillas es / pt", "code")]
            w, gap, x0, y0 = 1.72, 0.36, 0.6, 2.15
            for k, (name, sub, kind) in enumerate(nodes):
                x = x0 + k * (w + gap)
                shp = box(s, x, y0, w, 1.05, fill=BLUE if kind == "code" else WHITE,
                          line=None if kind == "code" else ORANGE)
                label_in(shp, f"{name}\n{sub}", size=15, color=WHITE if kind == "code" else INK)
                if k < len(nodes) - 1:
                    arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + w + 0.06), Inches(y0 + 0.4),
                                             Inches(gap - 0.12), Inches(0.25))
                    arr.fill.solid()
                    arr.fill.fore_color.rgb = INK_3
                    arr.line.fill.background()
            lg = box(s, 0.6, 3.4, 0.28, 0.2, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
            del lg
            text(s, 0.98, 3.35, 5, 0.3, "Código: decide y actúa (auditable, testeable)", size=13, color=INK_2)
            box(s, 6.3, 3.4, 0.28, 0.2, fill=WHITE, line=ORANGE, shape=MSO_SHAPE.RECTANGLE)
            text(s, 6.68, 3.35, 6, 0.3, "Modelo: interpreta (router aprendido local; si no carga, reglas)",
                 size=13, color=INK_2)
            for k, (name, desc) in enumerate(sd["layers"]):
                x = 0.6 + k * 3.08
                box(s, x, 4.15, 2.88, 2.2, fill=CARD)
                text(s, x + 0.22, 4.35, 2.5, 0.45, name, size=20, bold=True, color=INK)
                text(s, x + 0.22, 4.9, 2.5, 1.4, desc, size=15, color=INK_2)
        elif i == 3:
            for k, (name, msg, desc) in enumerate(sd["scenes"]):
                x = 0.6 + k * 4.1
                shot = os.path.join(PITCH, "screens", f"escena{k + 1}.png")
                if os.path.exists(shot):
                    image(s, shot, x, 2.05, w=3.85)
                else:
                    ph = box(s, x, 2.05, 3.85, 2.3, fill=CARD)
                    label_in(ph, f"[[PENDIENTE: captura]]\ndocs/pitch/screens/escena{k + 1}.png", size=13, color=PENDING)
                    if f"[[PENDIENTE: captura escena {k + 1} — docs/pitch/screens/escena{k + 1}.png (UI #14)]]" not in pending:
                        pending.append(f"[[PENDIENTE: captura escena {k + 1} — docs/pitch/screens/escena{k + 1}.png (UI #14)]]")
                text(s, x, 4.55, 3.85, 0.4, f"{k + 1} · {name}", size=20, bold=True, color=BLUE)
                text(s, x, 5.0, 3.85, 0.45, msg, size=15, italic=True, color=INK)
                text(s, x, 5.5, 3.85, 1.3, desc, size=13.5, color=INK_2)
        elif i == 4:
            for k, (name, big, sub, how) in enumerate(sd["metrics"]):
                x = 0.6 + k * 3.08
                box(s, x, 2.15, 2.88, 3.9, fill=TINT if k != 2 else CARD)
                text(s, x + 0.2, 2.35, 2.5, 0.4, name, size=15, bold=True, color=INK)
                text(s, x + 0.2, 2.85, 2.5, 1.2, big, size=30 if "[[" not in big else 15, bold=True, color=BLUE)
                text(s, x + 0.2, 4.05, 2.5, 0.6, sub, size=13, color=INK)
                text(s, x + 0.2, 4.75, 2.5, 1.2, how, size=12, color=INK_2)
            text(s, 0.6, 6.3, 12.1, 0.55, sd["foot"], size=14, color=INK)
            sd = dict(sd, foot="Detalle: reports/eval_router.md · eval.md · intent_router.md · injection_eval.md · redteam.md")
        elif i == 5:
            image(s, os.path.join(CHARTS, sd["chart"]), 0.5, 2.0, w=7.9)
            text(s, 8.75, 2.05, 4.0, 0.4, "Límite del dato → decisión", size=15, bold=True, color=INK_3)
            for k, (big, what, dec) in enumerate(sd["limits"]):
                y = 2.6 + k * 1.42
                text(s, 8.75, y, 1.3, 0.6, big, size=28, bold=True, color=BLUE)
                text(s, 10.05, y + 0.02, 2.75, 0.6, what, size=12.5, color=INK)
                text(s, 10.05, y + 0.62, 2.75, 0.6, f"→ {dec}", size=12.5, color=INK_2)
        elif i == 6:
            for k, (head, items, fill, mark) in enumerate([("Listo en el sandbox", sd["ready"], TINT, "✓"),
                                                          ("Falta para producción", sd["missing"], CARD, "○")]):
                x = 0.6 + k * 6.15
                box(s, x, 2.05, 5.95, 4.75, fill=fill)
                text(s, x + 0.3, 2.25, 5.4, 0.45, head, size=20, bold=True)
                text(s, x + 0.3, 2.9, 5.4, 3.8, [(t, {"space": 12, "mark": mark}) for t in items], size=16, color=INK)
        footer(s, sd["foot"])
        s.notes_slide.notes_text_frame.text = sd["notes"]
    prs.save(path)


# ------------------------------------------------------------------ render md
def render_md(slides, path):
    out = ["# Pitch · Agente de disputas de cargos no reconocidos", "",
           "> Generado por `scripts/make_deck.py` (misma fuente que `deck.pptx`). No editar a mano: editar el script o "
           "`docs/pitch/metrics.yaml` y regenerar.", ""]
    for i, sd in enumerate(slides, 1):
        out += [f"## {i}. {sd['title']}", f"*{sd['kicker']}*", ""]
        if "chart" in sd:
            out += [f"![{sd['chart']}](charts/{sd['chart']})", ""]
        for big, small in sd.get("tiles", []):
            out.append(f"- **{big}** {small}")
        for name, desc in sd.get("layers", []):
            out.append(f"- **{name}:** {desc.replace(chr(10), ' ')}")
        for k, (name, msg, desc) in enumerate(sd.get("scenes", []), 1):
            out.append(f"{k}. **{name}** — {msg}: {desc}")
        shots = [k for k in range(1, len(sd.get("scenes", [])) + 1)
                 if os.path.exists(os.path.join(PITCH, "screens", f"escena{k}.png"))]
        if shots:
            out += [""] + [f"![escena {k}](screens/escena{k}.png)" for k in shots]
        if "metrics" in sd:
            out += ["| Pregunta | Resultado | Contexto | Cómo se mide |", "|---|---|---|---|"]
            out += [f"| {a} | **{b}** | {c} | {d} |" for a, b, c, d in sd["metrics"]]
        for big, what, dec in sd.get("limits", []):
            out.append(f"- **{big}** {what} → {dec}")
        if "ready" in sd:
            out += ["**Listo en el sandbox**", ""] + [f"- {t}" for t in sd["ready"]]
            out += ["", "**Falta para producción**", ""] + [f"- {t}" for t in sd["missing"]]
        for e in sd.get("md_extra", []):
            out += ["", e]
        out += ["", f"<sub>{sd['foot']}</sub>", "", "**Notas del orador:** " + sd["notes"], "", "---", ""]
    if pending:
        out += ["## Pendientes", ""] + [f"- {p}" for p in pending] + [""]
    open(path, "w", encoding="utf-8").write("\n".join(out))


def main():
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    data, metrics, n_tests = load()
    slides = build(data, metrics, n_tests)
    render_pptx(slides, os.path.join(PITCH, "deck.pptx"))
    render_md(slides, os.path.join(PITCH, "deck.md"))
    print(f"docs/pitch/deck.pptx y docs/pitch/deck.md: {len(slides)} slides")
    if pending:
        print(f"{len(pending)} pendientes:")
        for p in pending:
            print("  " + p)


if __name__ == "__main__":
    main()
